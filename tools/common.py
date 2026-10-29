# -*- coding: utf-8 -*-
"""Utilidades compartidas del generador de talleres (solo se usa al construir el contenido)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR = PatternFill("solid", fgColor="1D4ED8")
SUB = PatternFill("solid", fgColor="E8EFFF")
THIN = Side(style="thin", color="C8D0E0")


def m(n):
    """Formato moneda colombiana: $1.234.567"""
    s = f"{abs(n):,.0f}".replace(",", ".")
    return ("-$" if n < 0 else "$") + s


def tmd(headers, rows):
    """Tabla markdown."""
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    for r in rows:
        out.append("| " + " | ".join(str(c) for c in r) + " |")
    return "\n".join(out)


def V(desc, value, puntos, sheet=None, label=None, tol=1, abs_=True, pista="", alts=None):
    d = {"tipo": "valor", "desc": desc, "value": value, "puntos": puntos, "tol": tol, "abs": abs_, "pista": pista}
    if sheet:
        d["sheet"] = sheet
    if label:
        d["label"] = label
    if alts:
        d["alts"] = alts
    return d


def F(sheet, minimo, puntos, desc):
    return {"tipo": "formulas", "sheet": sheet, "min": minimo, "puntos": puntos, "desc": desc,
            "pista": "Usa fórmulas de Excel (no valores escritos a mano) para que el modelo se actualice solo."}


def H(nombres, puntos, desc="Estructura del libro: hojas solicitadas"):
    return {"tipo": "hojas", "names": nombres, "puntos": puntos, "desc": desc,
            "pista": "Nombra las hojas exactamente como se pide en la sección «Qué debes entregar»."}


def R(criterio, puntos, detalle):
    return {"criterio": criterio, "puntos": puntos, "detalle": detalle}


def write_template(path, titulo, sheets):
    """sheets: lista de dicts {name, rows, widths, headers(list of row idx 1-based), note}"""
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh["name"][:31])
        for r in sh["rows"]:
            ws.append(r)
        for ri in sh.get("headers", []):
            for c in ws[ri]:
                if c.value is not None:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = HDR
                    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for ri in sh.get("subs", []):
            for c in ws[ri]:
                if c.value is not None:
                    c.font = Font(bold=True)
                    c.fill = SUB
        for i, w in enumerate(sh.get("widths", []), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.number_format = sh.get("numfmt", "#,##0")
                if isinstance(c.value, str) and len(c.value) > 60:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        if sh.get("freeze"):
            ws.freeze_panes = sh["freeze"]
        if sh.get("note"):
            ws.sheet_properties.tabColor = sh["note"]
    wb.properties.title = titulo
    wb.save(path)


def instr_sheet(titulo, lineas):
    rows = [[titulo], [""]] + [[l] for l in lineas]
    return {"name": "Instrucciones", "rows": rows, "widths": [110], "headers": [1], "note": "1D4ED8"}


def blank(name, headers, nrows=25, widths=None):
    rows = [headers] + [[None] * len(headers) for _ in range(nrows)]
    return {"name": name, "rows": rows, "headers": [1], "widths": widths or [18] * len(headers), "freeze": "A2", "note": "059669"}
